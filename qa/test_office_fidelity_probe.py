"""Synthetic, public-safe regression checks for the fidelity gate itself.

These checks do NOT verify an Android office layout engine or prove document fidelity.
"""
import tempfile
from pathlib import Path
import unittest
from zipfile import ZipFile
import fitz
from PIL import Image,ImageDraw
from docx import Document
from docx.shared import Inches
from office_fidelity_probe import audit,docx_structure,pdf_structure,compare_visual,MAX_DOC_XML_BYTES


def fixture(directory):
    root=Path(directory)
    img=root/"synthetic.png"
    bitmap=Image.new("RGB",(150,50),"white")
    painter=ImageDraw.Draw(bitmap)
    painter.ellipse((15,8,70,42),fill="green")
    bitmap.save(img)
    src=root/"synthetic.docx"
    d=Document()
    d.add_heading("Árvíztűrő tükörfúrógép",0)
    r=d.add_paragraph().add_run("Félkövér, színek, többféle font")
    r.bold=True
    t=d.add_table(rows=2,cols=2,style="Table Grid")
    t.cell(0,0).text="Táblázat"
    t.cell(1,0).text="Cella"
    d.add_picture(str(img),width=Inches(2))
    d.save(src)
    bad=root/"flattened.pdf"
    p=fitz.open()
    p.new_page().insert_text((40,50),"Text-only flattening")
    p.save(bad);p.close()
    reference=root/"reference.pdf"
    q=fitz.open()
    page=q.new_page()
    page.insert_text((40,50),"Office layout reference")
    page.insert_image(fitz.Rect(50,75,200,125),filename=str(img))
    q.save(reference);q.close()
    return src,bad,reference


class QualityProbeTests(unittest.TestCase):
    def test_detects_missing_docx_images(self):
        with tempfile.TemporaryDirectory() as temp:
            doc,bad,ref=fixture(temp)
            result=audit(doc,bad,ref)
            self.assertEqual(result["status"],"FAIL")
            self.assertIn("ALL_EMBEDDED_DOCX_MEDIA_MISSING",result["hard_failures"])
            self.assertEqual(result["source"]["table_nodes"],1)

    def test_identical_reference_does_not_mask_severe_source_text_loss(self):
        with tempfile.TemporaryDirectory() as temp:
            doc,bad,ref=fixture(temp)
            result=audit(doc,ref,ref)
            self.assertEqual(result["status"],"FAIL")
            self.assertIn("SEVERE_TEXT_LOSS",result["hard_failures"])
            self.assertEqual(result["visual"]["status"],"MEASURED_NOT_PASS")

    def test_good_text_coverage_is_review_required_not_pro_pass(self):
        with tempfile.TemporaryDirectory() as temp:
            folder=Path(temp)
            d=Document();d.add_paragraph("alpha beta gamma delta")
            doc=folder/"source.docx";d.save(doc)
            out=folder/"output.pdf"
            pdf=fitz.open();pdf.new_page().insert_text((40,50),"alpha beta gamma delta")
            pdf.save(out);pdf.close()
            result=audit(doc,out,out)
            self.assertEqual(result["status"],"REVIEW_REQUIRED")
            self.assertEqual(result["hard_failures"],[])
            self.assertEqual(result["visual"]["status"],"MEASURED_NOT_PASS")

    def test_invalid_docx_zip_rejected(self):
        with tempfile.TemporaryDirectory() as temp:
            invalid=Path(temp)/"broken.docx"
            invalid.write_bytes(b"not DOCX")
            with self.assertRaises(ValueError): docx_structure(invalid)

    def test_invalid_pdf_rejected(self):
        with tempfile.TemporaryDirectory() as temp:
            invalid=Path(temp)/"broken.pdf"
            invalid.write_bytes(b"not PDF")
            with self.assertRaises(Exception): pdf_structure(invalid)


    def test_severe_text_loss_is_hard_failure_even_without_images(self):
        with tempfile.TemporaryDirectory() as td:
            root=Path(td)
            d=Document();d.add_paragraph("Original formatted document with actual content "*12)
            src=root/"text.docx";d.save(src)
            p=fitz.open();p.new_page().insert_text((30,30),"Unrelated text")
            out=root/"lost.pdf";p.save(out);p.close()
            result=audit(src,out)
            self.assertEqual(result["status"],"FAIL")
            self.assertIn("SEVERE_TEXT_LOSS",result["hard_failures"])

    def test_partial_text_difference_requires_review_not_automatic_pass(self):
        with tempfile.TemporaryDirectory() as td:
            root=Path(td)
            d=Document();d.add_paragraph("alpha beta gamma delta epsilon zeta eta theta iota kappa")
            src=root/"text.docx";d.save(src)
            p=fitz.open();p.new_page().insert_text((30,30),"alpha beta gamma delta epsilon zeta eta")
            out=root/"partial.pdf";p.save(out);p.close()
            result=audit(src,out)
            self.assertEqual(result["status"],"REVIEW_REQUIRED")
            self.assertIn("LOW_TEXT_OVERLAP",result["review_warnings"])

    def test_docx_xml_with_dtd_is_rejected(self):
        with tempfile.TemporaryDirectory() as td:
            bad=Path(td)/"dtd.docx"
            with ZipFile(bad,"w") as z:
                z.writestr("word/document.xml", '<!DOCTYPE w:document [<!ENTITY x "unsafe">]><w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p>&x;</w:p></w:body></w:document>')
            with self.assertRaisesRegex(ValueError,"DTD/entity"):
                docx_structure(bad)

    def test_oversized_xml_denied_before_decompression(self):
        with tempfile.TemporaryDirectory() as td:
            bad=Path(td)/"huge.docx"
            with ZipFile(bad,"w",compression=8) as z:
                z.writestr("word/document.xml",b" "*(MAX_DOC_XML_BYTES+1))
            with self.assertRaisesRegex(ValueError,"document XML too large"):
                docx_structure(bad)

    def test_duplicate_document_parts_rejected(self):
        from warnings import catch_warnings, simplefilter
        with tempfile.TemporaryDirectory() as td:
            bad=Path(td)/"duplicate.docx"
            with catch_warnings():
                simplefilter("ignore",UserWarning)
                with ZipFile(bad,"w") as z:
                    z.writestr("word/document.xml",b"<xml/>")
                    z.writestr("word/document.xml",b"<xml/>")
            with self.assertRaisesRegex(ValueError,"duplicate document.xml"):
                docx_structure(bad)

    def test_huge_reference_page_is_not_rendered(self):
        with tempfile.TemporaryDirectory() as td:
            root=Path(td)
            for name in ("output.pdf","ref.pdf"):
                doc=fitz.open()
                doc.new_page(width=3000,height=3000).insert_text((10,20),"Hello")
                doc.save(root/name);doc.close()
            out=compare_visual(root/"output.pdf",root/"ref.pdf")
            self.assertEqual(out["status"],"MEASURED_NOT_PASS")
            self.assertEqual(out["pages"][0]["status"],"SKIPPED_RENDER_PIXEL_LIMIT")

if __name__=="__main__":
    unittest.main()
