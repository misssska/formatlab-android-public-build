"""Synthetic-only tests for the isolated LibreOffice reference converter.

Never put private customer DOCX/PDF files in this public test repository.
"""
import shutil
import tempfile
import unittest
from pathlib import Path
from libreoffice_reference import convert

class ReferenceConversionTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.directory = Path(self.temp.name)

    def make_docx(self):
        from docx import Document
        from docx.shared import Inches
        from PIL import Image
        img=self.directory/'test.png'
        Image.new('RGB',(30,30),'purple').save(img)
        src=self.directory/'synthetic.docx'
        document=Document()
        document.add_heading('Árvíztűrő tükörfúrógép',0)
        run=document.add_paragraph().add_run('Bold and italic')
        run.bold=True
        table=document.add_table(rows=2,cols=2,style='Table Grid')
        table.cell(0,0).text='Cell'
        document.add_picture(str(img),width=Inches(1.5))
        document.save(src)
        return src

    def test_rejects_non_office_input(self):
        src=self.directory/'test.txt'
        src.write_text('text',encoding='utf8')
        with self.assertRaisesRegex(ValueError,'Unsupported'):
            convert(src,self.directory/'out.pdf')

    def test_missing_local_engine_is_explicit(self):
        src=self.make_docx()
        with self.assertRaises(FileNotFoundError):
            convert(src,self.directory/'out.pdf',soffice='nonexistent_binary_formatlab_9876')

    @unittest.skipUnless(shutil.which('libreoffice'),'LibreOffice not installed')
    def test_real_offline_conversion_preserves_images_and_unicode(self):
        import fitz
        src=self.make_docx()
        output=self.directory/'converted.pdf'
        result=convert(src,output,timeout=180)
        self.assertEqual(result['status'],'REFERENCE_CREATED_NOT_ANDROID')
        self.assertGreaterEqual(result['embedded_images'],1)
        self.assertGreater(result['page_count'],0)
        with fitz.open(output) as pdf:
            text=' '.join(page.get_text() for page in pdf).lower()
            self.assertIn('tükörfúrógép',text)

if __name__=='__main__':unittest.main()
