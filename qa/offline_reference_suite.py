#!/usr/bin/env python3
"""Synthetic offline DOCX/PDF golden-reference generator; never tests Android itself."""
from __future__ import annotations
import argparse,hashlib,json,shutil,subprocess,sys,tempfile
from dataclasses import dataclass
from pathlib import Path
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm,Pt,RGBColor
from PIL import Image,ImageDraw
import fitz
from office_fidelity_probe import audit,docx_structure,pdf_structure

MAX_PDF_BYTES=30*1024*1024
MAX_DOCX_BYTES=10*1024*1024

@dataclass(frozen=True)
class Spec:
    name:str
    markers:tuple[str,...]
    minimum_pages:int
    minimum_images:int
    minimum_tables:int

SPECS=(
    Spec('styled_text',('FORMATLAB-STYLE-001','BOLD-ALPHA','BLUE-BETA'),1,0,0),
    Spec('tables_nested',('FORMATLAB-TABLE-002','MERGED-CELL','NESTED-CELL','LAST-ROW'),1,0,2),
    Spec('embedded_images',('FORMATLAB-IMAGE-003','IMAGE-BEFORE','IMAGE-AFTER'),1,3,0),
    Spec('pagination_unicode',('FORMATLAB-PAGES-004','FIRST-PAGE','SECOND-PAGE','THIRD-PAGE'),3,0,0),
    Spec('header_footer',('FORMATLAB-HF-005','HEADER-MARKER','FOOTER-MARKER','BODY-MARKER'),1,0,0),
)
def sha256(path:Path)->str:
    h=hashlib.sha256()
    with path.open('rb') as fp:
        for chunk in iter(lambda:fp.read(1024*1024),b''):h.update(chunk)
    return h.hexdigest()

def _png(path:Path,color:tuple[int,int,int],label:str):
    pic=Image.new('RGB',(480,190),color)
    draw=ImageDraw.Draw(pic)
    draw.rectangle((8,8,472,182),outline=(20,20,20),width=5)
    draw.ellipse((30,32,155,155),outline=(255,255,255),width=8)
    draw.text((185,76),label,fill=(255,255,255))
    pic.save(path,format='PNG',compress_level=6)

def _document(spec:Spec,images:Path)->Document:
    doc=Document()
    section=doc.sections[0]
    section.top_margin=section.bottom_margin=Cm(2)
    doc.styles['Normal'].font.name='Liberation Sans'
    doc.styles['Normal'].font.size=Pt(10)
    if spec.name=='styled_text':
        doc.add_heading('FORMATLAB-STYLE-001',0)
        p=doc.add_paragraph()
        p.add_run('BOLD-ALPHA').bold=True
        p.add_run(' Normal text, mixed ')
        span=p.add_run('BLUE-BETA');span.font.color.rgb=RGBColor(20,70,215)
        doc.add_paragraph('Hungarian vowels: árvíztűrő tükörfúrógép. 2026.')
        doc.add_paragraph('Numbers 12345 and punctuation: ; : ! ? .')
    elif spec.name=='tables_nested':
        doc.add_heading('FORMATLAB-TABLE-002',0)
        tab=doc.add_table(rows=4,cols=3)
        tab.alignment=WD_TABLE_ALIGNMENT.CENTER
        tab.style='Table Grid'
        tab.cell(0,0).text='MERGED-CELL'
        tab.cell(0,0).merge(tab.cell(0,1))
        tab.cell(0,2).text='HEADER'
        tab.cell(1,0).text='DATA-ONE'
        tab.cell(1,1).text='DATA-TWO'
        tab.cell(1,2).text='DATA-THREE'
        nested=tab.cell(2,1).add_table(rows=1,cols=1)
        nested.style='Table Grid'
        nested.cell(0,0).text='NESTED-CELL'
        tab.cell(3,0).text='LAST-ROW'
        doc.add_paragraph('After-table terminator.')
    elif spec.name=='embedded_images':
        doc.add_heading('FORMATLAB-IMAGE-003',0)
        doc.add_paragraph('IMAGE-BEFORE')
        for i,col in enumerate(((175,36,36),(35,155,55),(25,70,181)),1):
            path=images/f'sample-{i}.png'
            _png(path,col,f'IMG-{i}')
            doc.add_picture(str(path),width=Cm(7))
        doc.add_paragraph('IMAGE-AFTER')
    elif spec.name=='pagination_unicode':
        doc.add_heading('FORMATLAB-PAGES-004',0)
        doc.add_paragraph('FIRST-PAGE – Magyar áéíóöőúüű 你好')
        doc.add_page_break()
        doc.add_paragraph('SECOND-PAGE – český; русский; العربية')
        doc.add_page_break()
        doc.add_paragraph('THIRD-PAGE – Ελληνικά; français; español')
    elif spec.name=='header_footer':
        section.header.paragraphs[0].text='HEADER-MARKER'
        footer=section.footer.paragraphs[0]
        footer.text='FOOTER-MARKER'
        footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
        doc.add_heading('FORMATLAB-HF-005',0)
        doc.add_paragraph('BODY-MARKER with running section style.')
    else:raise ValueError('Unknown fixture')
    return doc

def create_corpus(root:Path)->list[tuple[Spec,Path]]:
    root=Path(root);root.mkdir(parents=True,exist_ok=True)
    image_dir=root/'generated-images';image_dir.mkdir(exist_ok=True)
    corpus=[]
    for spec in SPECS:
        path=root/f'{spec.name}.docx'
        _document(spec,image_dir).save(path)
        if path.stat().st_size>MAX_DOCX_BYTES:raise ValueError('GENERATED_DOCX_TOO_LARGE')
        corpus.append((spec,path))
    return corpus

def convert_reference(src:Path,output_dir:Path,soffice:str,*,timeout:int=90)->Path:
    if not 5<=timeout<=300:raise ValueError('INVALID_TIMEOUT')
    output_dir=Path(output_dir);output_dir.mkdir(parents=True,exist_ok=True)
    with tempfile.TemporaryDirectory(prefix='formatlab-soffice-profile-') as tmp:
        profile=(Path(tmp)/'profile').as_uri()
        try:
            run=subprocess.run([soffice,f'-env:UserInstallation={profile}',
                '--headless','--convert-to','pdf:writer_pdf_Export',
                '--outdir',str(output_dir),str(src)],timeout=timeout,
                stdout=subprocess.PIPE,stderr=subprocess.PIPE,text=True,check=False)
        except (OSError,subprocess.TimeoutExpired) as exc:
            raise RuntimeError('LIBREOFFICE_CONVERSION_FAILED_OR_TIMED_OUT') from exc
    pdf=output_dir/(src.stem+'.pdf')
    if run.returncode or not pdf.is_file() or pdf.stat().st_size<100:
        raise RuntimeError('LIBREOFFICE_REFERENCE_MISSING: '+run.stderr[-500:])
    if pdf.stat().st_size>MAX_PDF_BYTES:raise RuntimeError('REFERENCE_PDF_TOO_LARGE')
    return pdf

def evaluate(spec:Spec,docx:Path,pdf:Path)->dict:
    src=docx_structure(docx);out=pdf_structure(pdf)
    with fitz.open(pdf) as pages:
        text='\n'.join(p.get_text(sort=True) for p in pages)
        count=len(pages)
    compact=''.join(text.split()).casefold()
    missing=[m for m in spec.markers if ''.join(m.split()).casefold() not in compact]
    issues=[]
    if missing:issues.append('REQUIRED_TEXT_MARKERS_MISSING')
    if count<spec.minimum_pages:issues.append('EXPECTED_PAGE_BREAKS_MISSING')
    if out['embedded_images']<spec.minimum_images:issues.append('EMBEDDED_MEDIA_MISSING')
    if src['table_nodes']<spec.minimum_tables:issues.append('BAD_SOURCE_FIXTURE_TABLE_COUNT')
    if src['media_files']<spec.minimum_images:issues.append('BAD_SOURCE_FIXTURE_MEDIA_COUNT')
    check=audit(docx,pdf)
    if check['status']=='FAIL':issues.append('SOURCE_TO_PDF_HARD_FAILURE')
    return {'case':spec.name,'result':'REFERENCE_PASS' if not issues else 'FAIL',
            'issues':issues,'missing_markers':missing,'pages':count,
            'embedded_images':out['embedded_images'],'table_nodes':src['table_nodes'],
            'word_overlap':check['word_overlap'],'docx_sha256':sha256(docx),
            'pdf_sha256':sha256(pdf),'visual_fidelity_compared_to_android':False}

def build(root:Path,*,soffice:str,timeout:int=90)->dict:
    root=Path(root).resolve()
    reference=root/'reference_pdf';corpus=root/'synthetic_docx'
    results=[]
    for spec,src in create_corpus(corpus):
        pdf=convert_reference(src,reference,soffice,timeout=timeout)
        results.append(evaluate(spec,src,pdf))
    report={'schema_version':1,'suite':'formatlab-offline-office-golden',
            'source_data':'synthetic','cases':results,
            'status':'REFERENCE_PASS' if all(x['result']=='REFERENCE_PASS' for x in results) else 'FAIL',
            'android_converted':False,'android_fidelity_tested':False,
            'notes':'LibreOffice Linux reference only. Android output must be checked separately.'}
    (root/'suite-report.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    return report

def main(argv=None)->int:
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument('--output',type=Path,required=True)
    p.add_argument('--soffice',default=shutil.which('libreoffice') or 'libreoffice')
    p.add_argument('--timeout',type=int,default=90)
    a=p.parse_args(argv)
    try:
        result=build(a.output,soffice=a.soffice,timeout=a.timeout)
        print(json.dumps({'status':result['status'],
            'cases':[(x['case'],x['result'],x['pages'],x['embedded_images']) for x in result['cases']]},ensure_ascii=False))
        return 0 if result['status']=='REFERENCE_PASS' else 2
    except (OSError,ValueError,RuntimeError) as exc:
        print('OFFLINE_REFERENCE_SUITE_FAILED:',str(exc),file=sys.stderr)
        return 3

if __name__=='__main__':raise SystemExit(main())
